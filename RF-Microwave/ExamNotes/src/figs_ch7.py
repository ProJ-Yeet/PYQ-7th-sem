# -*- coding: utf-8 -*-
"""Chapter 7 figures, lifted from the sources and never redrawn.

Original embedded image blobs from Er. Kobid Karkee's Chapter_7 PDF, the
KEC Chapter_7_New deck, Er. Shankar Gangaju's Chapter_7 deck and a 2078
student presentation. A job names the source, the slide (pptx) or page (pdf,
1-based), which picture on it, an optional crop box as fractions
(left, top, right, bottom), and whether to invert (the PDF's zone diagram is
white line art on black, unreadable in print).

Run from src\\:   python figs_ch7.py
"""
import io
import os

import fitz
from PIL import Image, ImageOps
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
RP = os.path.normpath(os.path.join(HERE, "..", "..", "Notes", "RF Pulchowk"))
STU = os.path.join(RP, "New_Chapter_5", "students presentation_2078")
SRC = {
    "pdf": os.path.join(RP, "Chapter 7", "Chapter_7.pdf"),
    "gangaju": os.path.join(RP, "all", "Chapter_7.pptx"),
    "kec": os.path.join(RP, "all", "Chapter_7_New.pptx"),
    "ant": os.path.join(STU, "Microwave antenna and Signal characteristics.pptx"),
}
OUT = os.path.join(HERE, "figs")

# name: (source, slide or page, picture index, crop box or None, invert)
JOBS = {
    "c7_zones":    ("pdf", 5, 6, None, True),
    "c7_fcc":      ("pdf", 18, 6, None, False),
    "c7_ansi":     ("pdf", 19, 6, (0.07, 0.235, 1.0, 0.95), False),
    "c7_heating":  ("gangaju", 4, 0, None, False),
    "c7_spectrum": ("kec", 5, 0, None, False),
    "c7_warning":  ("kec", 24, 0, (0.0, 0.0, 1.0, 0.585), False),
    "c7_tower":    ("kec", 32, 0, (0.0, 0.135, 1.0, 0.68), False),
    "c7_dish":     ("ant", 6, 0, None, False),
    "c7_horn":     ("ant", 7, 0, None, False),
    "c7_patch":    ("ant", 8, 0, None, False),
}


def flatten(im):
    if im.mode in ("RGBA", "LA", "P"):
        im = im.convert("RGBA")
        bg = Image.new("RGB", im.size, "white")
        bg.paste(im, mask=im.split()[-1])
        return bg
    return im.convert("RGB")


def pptx_pic(path, slide, idx):
    prs = Presentation(path)
    pics = [Image.open(io.BytesIO(sh.image.blob))
            for sh in prs.slides[slide - 1].shapes if sh.shape_type == 13]
    return flatten(pics[idx])


def pdf_pic(path, page, idx):
    doc = fitz.open(path)
    # raw image index: every page of this export carries six background
    # images first, so the figure is index 6
    info = doc[page - 1].get_images(full=True)[idx]
    pm = fitz.Pixmap(doc, info[0])
    if pm.n - pm.alpha > 3:
        pm = fitz.Pixmap(fitz.csRGB, pm)
    return flatten(Image.open(io.BytesIO(pm.tobytes("png"))))


def main():
    os.makedirs(OUT, exist_ok=True)
    for name, (src, n, idx, box, inv) in JOBS.items():
        path = SRC[src]
        im = pdf_pic(path, n, idx) if src == "pdf" else pptx_pic(path, n, idx)
        if box:
            w, h = im.size
            im = im.crop((int(box[0] * w), int(box[1] * h), int(box[2] * w), int(box[3] * h)))
        if inv:
            im = ImageOps.invert(im)
        g = im.convert("L")
        dark = sum(1 for v in g.get_flattened_data() if v < 200) / float(g.size[0] * g.size[1])
        if dark < 0.005 or dark > 0.9:
            raise AssertionError("%s looks blank (ink %.3f)" % (name, dark))
        im.save(os.path.join(OUT, name + ".png"))
        print("%-12s %-8s %2d  %4dx%-4d ink %.3f" % (name, src, n, im.size[0], im.size[1], dark))


if __name__ == "__main__":
    main()
