# -*- coding: utf-8 -*-
"""Chapter 6 source figures, lifted from Er. Gangaju's Chapter_6 deck.

Every slide body in this deck is one embedded picture, so a figure is a
fractional crop (left, top, right, bottom) of that original blob. The
microstrip layout drawings are not here: they are generated to scale from
computed dimensions by filt.py.

Run from src\\:   python figs_ch6.py
"""
import io
import os

from PIL import Image
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
DECK = os.path.normpath(os.path.join(HERE, "..", "..", "Notes", "RF Pulchowk", "all", "Chapter_6.pptx"))
OUT = os.path.join(HERE, "figs")

JOBS = {
    "c6_ladder":       (7, (0.00, 0.00, 1.00, 0.90)),
    "c6_responses":    (11, None),
    "c6_transform":    (19, (0.00, 0.00, 0.79, 0.92)),
    "c6_stepped_deck": (26, None),
    "c6_elliptic":     (27, None),
    "c6_stab_circles": (46, (0.44, 0.00, 1.00, 1.00)),
    "c6_gain_circles": (55, (0.46, 0.00, 1.00, 1.00)),
    "c6_amp_final":    (58, None),
    "c6_feedback":     (60, (0.00, 0.00, 0.47, 0.42)),
    "c6_twoport_osc":  (62, (0.00, 0.00, 0.93, 0.46)),
    "c6_oneport":      (65, (0.595, 0.00, 1.00, 0.41)),
    "c6_osc_steps":    (67, (0.60, 0.00, 1.00, 0.60)),
}


def main():
    os.makedirs(OUT, exist_ok=True)
    prs = Presentation(DECK)
    for name, (sl, box) in JOBS.items():
        pic = [s for s in prs.slides[sl - 1].shapes if s.shape_type == 13 and s.width > prs.slide_width * 0.3][0]
        im = Image.open(io.BytesIO(pic.image.blob)).convert("RGB")
        if box:
            w, h = im.size
            im = im.crop((int(box[0] * w), int(box[1] * h), int(box[2] * w), int(box[3] * h)))
        g = im.convert("L")
        dark = sum(1 for v in g.get_flattened_data() if v < 200) / float(g.size[0] * g.size[1])
        if dark < 0.005:
            raise AssertionError(name + " looks blank")
        im.save(os.path.join(OUT, name + ".png"))
        print("%-16s slide %2d %4dx%-4d ink %.3f" % (name, sl, im.size[0], im.size[1], dark))


if __name__ == "__main__":
    main()
