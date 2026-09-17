# -*- coding: utf-8 -*-
"""Chapter 8 figures, lifted from the sources and never redrawn.

Original embedded image blobs from Er. Kobid Karkee's Chapter_8 deck. Note
the deck (`all/Chapter_8.pptx`, 39 slides) is a SUPERSET of its own PDF export
(`Chapter 8/Chapter_8.pdf`, 28 pages): the impedance-measurement, reflectometer
and low/high-VSWR slides (23-30) are missing from the PDF entirely, and the
high-VSWR slide is the only source drawing of the double-minimum method. Pull
from the pptx.

A job names the slide, which picture on it, an optional crop box as fractions
(left, top, right, bottom), and whether to invert.

Run from src\\:   python figs_ch8.py            (add --sheet for a contact sheet)
"""
import io
import os
import sys

from PIL import Image, ImageOps
from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
DECK = os.path.normpath(os.path.join(
    HERE, "..", "..", "Notes", "RF Pulchowk", "all", "Chapter_8.pptx"))
OUT = os.path.join(HERE, "figs")

# photographs, not line art: these are legitimately dark, so the blank test
# only checks they are not a solid block of one tone
PHOTOS = {"c8_lowvswr2"}

# name: (slide, picture index, crop box or None, invert)
JOBS = {
    "c8_schottky":      (6, 0, None, False),
    "c8_barretter":     (8, 0, None, False),
    "c8_bridge":        (9, 0, None, False),
    "c8_singlebridge":  (10, 0, None, False),
    "c8_doublebridge":  (12, 0, None, False),
    "c8_thermocouple":  (14, 0, None, False),
    # slide 16 carries the same drawing as slide 19; kept once
    "c8_static":        (18, 0, None, False),
    "c8_circulating":   (19, 0, None, False),
    "c8_wattmeter":     (20, 0, None, False),
    "c8_slotted":       (22, 0, None, False),
    "c8_impedance":     (24, 0, None, False),
    "c8_reflectometer": (26, 0, None, False),
    "c8_vswrsetup":     (27, 0, None, False),
    "c8_lowvswr":       (28, 0, None, False),
    "c8_lowvswr2":      (28, 1, None, False),
    # slide 29 picture 0 is a solid colour block, not a drawing; picture 1 is
    # the only source drawing of the double-minimum construction anywhere
    "c8_highvswr":      (29, 1, None, False),
    "c8_vswrmeter":     (31, 0, None, False),
    "c8_spectrum":      (33, 0, None, False),
    "c8_vna":           (36, 0, None, False),
}


def flatten(im):
    if im.mode in ("RGBA", "LA", "P"):
        im = im.convert("RGBA")
        bg = Image.new("RGB", im.size, "white")
        bg.paste(im, mask=im.split()[-1])
        return bg
    return im.convert("RGB")


def pptx_pic(prs, slide, idx):
    pics = [Image.open(io.BytesIO(sh.image.blob))
            for sh in prs.slides[slide - 1].shapes if sh.shape_type == 13]
    return flatten(pics[idx])


def ink(im):
    g = im.convert("L")
    return sum(1 for v in g.get_flattened_data() if v < 200) / float(g.size[0] * g.size[1])


def sheet(images):
    """One contact sheet, 4 across, so the whole set can be judged in one look."""
    cols, cw, ch = 4, 340, 280
    rows = (len(images) + cols - 1) // cols
    out = Image.new("RGB", (cols * cw, rows * ch), "white")
    for k, (name, im) in enumerate(images):
        t = im.copy()
        t.thumbnail((cw - 12, ch - 26))
        x = (k % cols) * cw + (cw - t.size[0]) // 2
        y = (k // cols) * ch + 20
        out.paste(t, (x, y))
    out.save(os.path.join(HERE, "figs_ch8_sheet.png"))
    print("contact sheet -> figs_ch8_sheet.png  (order: " +
          ", ".join(n for n, _ in images) + ")")


def main():
    os.makedirs(OUT, exist_ok=True)
    prs = Presentation(DECK)
    made = []
    for name, (n, idx, box, inv) in JOBS.items():
        im = pptx_pic(prs, n, idx)
        if box:
            w, h = im.size
            im = im.crop((int(box[0] * w), int(box[1] * h),
                          int(box[2] * w), int(box[3] * h)))
        if inv:
            im = ImageOps.invert(im)
        d = ink(im)
        hi = 0.995 if name in PHOTOS else 0.9
        if d < 0.005 or d > hi:
            raise AssertionError("%s looks blank (ink %.3f)" % (name, d))
        im.save(os.path.join(OUT, name + ".png"))
        made.append((name, im))
        print("%-18s slide %2d  %4dx%-4d ink %.3f" % (name, n, im.size[0], im.size[1], d))
    if "--sheet" in sys.argv:
        sheet(made)


if __name__ == "__main__":
    main()
